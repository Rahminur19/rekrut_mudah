import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6] # completely blank layout

    # Color Palette
    BG_DARK = RGBColor(15, 23, 42)       # Slate 900
    BG_LIGHT = RGBColor(248, 250, 252)   # Slate 50
    CARD_BG = RGBColor(255, 255, 255)    # Pure White
    CARD_BORDER = RGBColor(226, 232, 240)# Slate 200
    PRIMARY = RGBColor(67, 56, 202)      # Indigo 700
    PRIMARY_LIGHT = RGBColor(99, 102, 241)# Indigo 500
    ACCENT_CYAN = RGBColor(14, 165, 233) # Sky 500
    ACCENT_TEAL = RGBColor(13, 148, 136) # Teal 600
    ACCENT_AMBER = RGBColor(245, 158, 11)# Amber 500
    TEXT_DARK = RGBColor(30, 41, 59)     # Slate 800
    TEXT_MUTED = RGBColor(100, 116, 139) # Slate 500
    TEXT_WHITE = RGBColor(255, 255, 255) # White

    def add_header(slide, title_text, category_text="REKRUT MUDAH • PRESENTASI SISTEM"):
        # Header category badge
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(10), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PRIMARY_LIGHT
        p_cat.font.name = 'Arial'

        # Main slide title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_DARK
        p_title.font.name = 'Arial'

    # ==========================================
    # SLIDE 1: COVER (DARK THEME)
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = BG_DARK
    bg1.line.fill.background()

    # Decorative accent bar
    bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), Inches(0.12), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_CYAN
    bar.line.fill.background()

    # Badge pill
    badge = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.45), Inches(4.5), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = RGBColor(30, 41, 59)
    badge.line.color.rgb = PRIMARY_LIGHT
    tf_b = badge.text_frame
    p_b = tf_b.paragraphs[0]
    p_b.text = "PROPOSAL & PERANCANGAN SISTEM v1.0"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_CYAN
    p_b.alignment = PP_ALIGN.CENTER

    # Main Title
    tbox = slide1.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(11), Inches(1.4))
    tf = tbox.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "REKRUT MUDAH"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_WHITE
    p1.font.name = 'Arial'

    p2 = tf.add_paragraph()
    p2.text = "Platform Rekrutmen Vokasi, Magang Mahasiswa & Penempatan Siswa PKL SMK"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(148, 163, 184)
    p2.font.name = 'Arial'

    # Tech Stack & Highlight Pills
    tbox2 = slide1.shapes.add_textbox(Inches(1.2), Inches(3.8), Inches(11), Inches(1.0))
    tf2 = tbox2.text_frame
    p3 = tf2.paragraphs[0]
    p3.text = "Arsitektur DATH Stack: Django 5 • Alpine.js 3 • Tailwind CSS • HTMX 2"
    p3.font.size = Pt(14)
    p3.font.bold = True
    p3.font.color.rgb = PRIMARY_LIGHT

    p4 = tf2.add_paragraph()
    p4.text = "Design System: Google Material Design 3 (M3) | Modern Monolith / HTML-over-the-Wire"
    p4.font.size = Pt(12)
    p4.font.color.rgb = RGBColor(148, 163, 184)

    # Footer Card on Slide 1
    f_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.6), Inches(10.9), Inches(1.1))
    f_card.fill.solid()
    f_card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    f_card.line.color.rgb = RGBColor(51, 65, 85)
    f_tf = f_card.text_frame
    f_p = f_tf.paragraphs[0]
    f_p.text = "Daftar Isi Presentasi: 1. PRD & Masalah  •  2. ERD & Database  •  3. Arsitektur DATH  •  4. Fitur Utama & ATS  •  5. Roadmap"
    f_p.font.size = Pt(12)
    f_p.font.color.rgb = RGBColor(226, 232, 240)
    f_p.alignment = PP_ALIGN.CENTER

    # Helper function for card
    def add_card(slide, x, y, w, h, title, items, badge_text=None, badge_color=PRIMARY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        offset_y = y + 0.2
        if badge_text:
            bdg = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.3), Inches(offset_y), Inches(len(badge_text)*0.13 + 0.4), Inches(0.32))
            bdg.fill.solid()
            bdg.fill.fore_color.rgb = badge_color
            bdg.line.fill.background()
            p_b = bdg.text_frame.paragraphs[0]
            p_b.text = badge_text
            p_b.font.size = Pt(9)
            p_b.font.bold = True
            p_b.font.color.rgb = TEXT_WHITE
            p_b.alignment = PP_ALIGN.CENTER
            offset_y += 0.42

        # Card Title
        tb_title = slide.shapes.add_textbox(Inches(x + 0.25), Inches(offset_y), Inches(w - 0.5), Inches(0.45))
        p_t = tb_title.text_frame.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_DARK

        # Card Body
        tb_body = slide.shapes.add_textbox(Inches(x + 0.25), Inches(offset_y + 0.45), Inches(w - 0.5), Inches(h - (offset_y - y + 0.5)))
        tf_body = tb_body.text_frame
        tf_body.word_wrap = True
        for i, itm in enumerate(items):
            p = tf_body.paragraphs[0] if i == 0 else tf_body.add_paragraph()
            p.text = f"• {itm}"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_DARK if not itm.startswith("Ket:") else TEXT_MUTED
            p.space_after = Pt(4)

    # ==========================================
    # SLIDE 2: LATAR BELAKANG & PERMASALAHAN (PRD 1.1)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "1. Latar Belakang & Masalah di Lapangan (PRD Section 1)")
    
    col_w = 3.65
    add_card(slide2, 0.8, 1.6, col_w, 5.2, "Tantangan Siswa SMK (PKL)", [
        "Mencari tempat PKL mandiri secara konvensional (door-to-door).",
        "Sering ditolak karena perusahaan tidak paham kompetensi kurikulum SMK.",
        "Kebutuhan unggah berkas resmi (Surat Pengantar Sekolah) tidak diakomodasi portal umum.",
        "Ketidakpastian durasi penempatan (3 bulan, 6 bulan, atau 1 tahun)."
    ], "MASALAH SISWA", RGBColor(225, 29, 72))

    add_card(slide2, 4.8, 1.6, col_w, 5.2, "Pencari Kerja & Mahasiswa", [
        "Minim transparansi benefit, fasilitas kerja, dan uang saku magang.",
        "Proses lamaran panjang (multi-step form) yang membosankan.",
        "Status lamaran 'menggantung' tanpa kejelasan feedback.",
        "Tidak ada informasi visual fasilitas lab/lingkungan industri."
    ], "MASALAH PELAMAR", RGBColor(217, 119, 6))

    add_card(slide2, 8.8, 1.6, col_w, 5.2, "Perekrut & HRD Industri", [
        "Kewalahan mengelola berkas lamaran yang tercecer via email.",
        "Tidak memiliki Applicant Tracking System (ATS) yang terstruktur.",
        "Sulit memverifikasi kompetensi siswa sesuai jurusan SMK (RPL, TKJ, dll.).",
        "Ketiadaan kanal komunikasi instan sebelum proses interview formal."
    ], "MASALAH INDUSTRI", RGBColor(14, 165, 233))

    # ==========================================
    # SLIDE 3: SOLUSI PRODUK & 3 PILAR KARIR (PRD 1.2)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "2. Solusi Rekrut Mudah: 3 Pilar Peluang Karir Terpadu")

    add_card(slide3, 0.8, 1.6, col_w, 5.2, "🎓 Pilar PKL SMK & Vokasi", [
        "Filter spesifik kompetensi keahlian: RPL, TKJ, Otomotif, AKL, OTKP.",
        "Dukungan unggah berkas resmi Surat Pengantar PKL dari sekolah.",
        "Filter durasi fleksibel (3 bulan, 6 bulan, 12 bulan).",
        "Menampilkan fasilitas lab/bengkel praktik industri langsung di profil perusahaan."
    ], "FOKUS VOKASI", ACCENT_TEAL)

    add_card(slide3, 4.8, 1.6, col_w, 5.2, "💼 Pilar Magang Mahasiswa", [
        "Program Internship bersertifikat dan berbayar.",
        "Transparansi benefit: uang saku, mentoring 1-on-1, konversi SKS.",
        "One-click apply instan menggunakan profil yang tersimpan.",
        "Pemberitahuan real-time tahapan seleksi (Review, Interview, Offer)."
    ], "INTERNSHIP", PRIMARY)

    add_card(slide3, 8.8, 1.6, col_w, 5.2, "🏢 Pilar Karir Profesional", [
        "Peluang kerja Full-time, Part-time, Contract, dan Remote.",
        "Rentang gaji transparan & filter sistem kerja (Onsite/Hybrid/Remote).",
        "Fitur Bookmark & Q&A diskusi publik pada setiap lowongan.",
        "Reputasi perusahaan dengan rating bintang terverifikasi (1.0 - 5.0)."
    ], "PROFESIONAL", ACCENT_AMBER)

    # ==========================================
    # SLIDE 4: TARGET PERSONA PENGGUNA (PRD Section 2)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "3. Target Persona Pengguna (User Personas)")

    p_w = 5.7
    add_card(slide4, 0.8, 1.6, p_w, 2.5, "1. Siswa SMK (PKL Seeker)", [
        "Profil: Ahmad (17 th), Siswa RPL SMKN 1 Jakarta.",
        "Perilaku: Cepat, mobile-first via smartphone atau laptop sekolah.",
        "Kebutuhan: Cari kuota PKL valid, upload Surat Pengantar, live chat HRD."
    ], "SISWA PKL", ACCENT_TEAL)

    add_card(slide4, 6.8, 1.6, p_w, 2.5, "2. Mahasiswa / Fresh Graduate", [
        "Profil: Budi (22 th), Sarjana Komputer / Fresh Graduate.",
        "Perilaku: Meneliti detail gaji, benefit, dan budaya kerja perusahaan.",
        "Kebutuhan: Lamar posisi Junior Developer, pantau histori seleksi di ATS."
    ], "JOB SEEKER", PRIMARY)

    add_card(slide4, 0.8, 4.3, p_w, 2.5, "3. Perekrut / HRD Perusahaan", [
        "Profil: Dewi Lestari, Talent Acquisition PT Nusantara Teknologi Mandiri.",
        "Perilaku: Mengelola ratusan berkas masuk setiap minggu.",
        "Kebutuhan: Pasang loker, seleksi ATS tanpa reload, jadwalkan wawancara."
    ], "RECRUITER", RGBColor(14, 165, 233))

    add_card(slide4, 6.8, 4.3, p_w, 2.5, "4. Superadministrator Sistem", [
        "Profil: Tim Verifikator Rekrut Mudah.",
        "Perilaku: Memastikan keamanan platform dan integritas mitra industri.",
        "Kebutuhan: Verifikasi legalitas perusahaan, kelola master jurusan & moderasi konten."
    ], "SUPERADMIN", RGBColor(71, 85, 105))

    # ==========================================
    # SLIDE 5: KEBUTUHAN FUNGSIONAL (PRD FR-01 s/d FR-06)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "4. Spesifikasi Kebutuhan Fungsional Produk (PRD Section 3)")

    f_w = 3.65
    add_card(slide5, 0.8, 1.6, f_w, 2.5, "FR-01: Multi-Role Auth & RBAC", [
        "Role: Siswa PKL, Pelamar Umum, HRD Perusahaan.",
        "Profil Khusus: Dokumen CV & Surat Pengantar PKL.",
        "Profil PT: Banner, logo, fasilitas & izin PKL."
    ], "MODUL 1", PRIMARY)

    add_card(slide5, 4.8, 1.6, f_w, 2.5, "FR-02: Live Search & Filter (HTMX)", [
        "Search-as-you-type debounced 300ms.",
        "Filter multivariabel instan (Tipe, Onsite/Remote, Jurusan).",
        "Bookmark & bookmark count update tanpa reload."
    ], "MODUL 2", ACCENT_CYAN)

    add_card(slide5, 8.8, 1.6, f_w, 2.5, "FR-03: One-Click Apply Modal", [
        "Dialog modal dimuat on-demand via HTMX.",
        "Pilihan berkas tersimpan atau upload baru.",
        "Pencegahan lamaran duplikat (Unique constraint)."
    ], "MODUL 3", ACCENT_TEAL)

    add_card(slide5, 0.8, 4.3, f_w, 2.5, "FR-04: Chat In-App Interaktif", [
        "Tombol 'Tanya HRD' di detail loker & profil PT.",
        "Live stream percakapan dengan background polling 4s.",
        "M3 Chat bubbles (Primary Container vs Surface High)."
    ], "MODUL 4", RGBColor(225, 29, 72))

    add_card(slide5, 4.8, 4.3, f_w, 2.5, "FR-05: Kanban ATS & Screening", [
        "Pipeline seleksi: Terkirim, Review, Interview, Terima, Tolak.",
        "Ubah status kandidat seketika via dropdown HTMX.",
        "Input jadwal & tautan wawancara interaktif."
    ], "MODUL 5", ACCENT_AMBER)

    add_card(slide5, 8.8, 4.3, f_w, 2.5, "FR-06: Profil Perusahaan & Rating", [
        "Rating Bintang Industri (1.0 s/d 5.0) & Ulasan.",
        "Galeri foto interaktif fasilitas lab & suasana kerja.",
        "Status badge verified dan kebijakan penerimaan PKL."
    ], "MODUL 6", RGBColor(124, 58, 237))

    # ==========================================
    # SLIDE 6: ARSITEKTUR TEKNOLOGI (DATH STACK)
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "5. Arsitektur Sistem: DATH Stack & Modern Monolith")

    # Stack components in 4 horizontal cards
    s_w = 2.7
    add_card(slide6, 0.8, 1.6, s_w, 4.0, "Django 5.x", [
        "Core Backend & ORM",
        "UUID primary keys untuk keamanan id.",
        "Role-Based Access Control (RBAC).",
        "Django Auth, CSRF & XSS protection.",
        "SQLite / PostgreSQL support."
    ], "BACKEND", RGBColor(22, 101, 52))

    add_card(slide6, 3.8, 1.6, s_w, 4.0, "Alpine.js 3.x", [
        "Micro-Interactions Client",
        "Dropdowns & Navigation drawers.",
        "Global Toast Notification store.",
        "Auto-scroll chat messages.",
        "Zero-build overhead (< 15KB)."
    ], "CLIENT SCRIPT", RGBColor(14, 116, 144))

    add_card(slide6, 6.8, 1.6, s_w, 4.0, "Tailwind CSS", [
        "Utility-First Styling",
        "Responsive Grid & Flexbox.",
        "M3 Design Tokens & Elevations.",
        "Dark/Light adaptive palette.",
        "Rapid component assembly."
    ], "UI STYLING", RGBColor(3, 105, 161))

    add_card(slide6, 9.8, 1.6, s_w, 4.0, "HTMX 2.x", [
        "HTML-over-the-Wire",
        "AJAX tanpa boilerplate JS.",
        "Live Search & Filter instan.",
        "Partial DOM replacement.",
        "Server-rendered HTML partials."
    ], "DYNAMIC DATA", RGBColor(194, 65, 12))

    # Architecture banner bottom
    b_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0))
    b_card.fill.solid()
    b_card.fill.fore_color.rgb = RGBColor(241, 245, 249)
    b_card.line.color.rgb = CARD_BORDER
    b_tf = b_card.text_frame
    b_p = b_tf.paragraphs[0]
    b_p.text = "Keunggulan DATH Stack: Memberikan performa interaktif setara SPA (Single Page Application) tanpa kompleksitas build-step berat, pemeliharaan REST API ganda, ataupun overhead bundle JavaScript."
    b_p.font.size = Pt(11)
    b_p.font.color.rgb = TEXT_DARK
    b_p.alignment = PP_ALIGN.CENTER

    # ==========================================
    # SLIDE 7: ENTITY RELATIONSHIP DIAGRAM (ERD)
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "6. Entity Relationship Diagram (ERD) & Struktur Database")

    erd_w = 3.65
    add_card(slide7, 0.8, 1.6, erd_w, 2.5, "Entitas Pengguna (Accounts)", [
        "User: UUID (PK), email (UK), role, phone, is_verified.",
        "UserProfile: user_id (FK), candidate_type, education_level, resume_file, recommendation_letter_file (Surat PKL).",
        "CompanyProfile: user_id (FK), name, slug, rating, accepts_pkl."
    ], "AKUN & PROFIL", PRIMARY)

    add_card(slide7, 4.8, 1.6, erd_w, 2.5, "Entitas Lowongan (Jobs)", [
        "JobCategory: id (PK), name, slug, icon_name.",
        "Major (Jurusan SMK): id (PK), name, level (SMK/D3/S1).",
        "JobListing: UUID (PK), company_id (FK), category_id (FK), title, job_type, majors (M2M), stipend, deadline."
    ], "LOWONGAN & JURUSAN", ACCENT_TEAL)

    add_card(slide7, 8.8, 1.6, erd_w, 2.5, "Entitas Interaksi (Jobs)", [
        "JobBookmark: user_id (FK), job_id (FK) -> Wishlist.",
        "JobDiscussion: job_id (FK), user_id (FK), question, reply, replied_by.",
        "CompanyGalleryPhoto: company_id (FK), photo, caption."
    ], "INTERAKSI & MEDIA", ACCENT_AMBER)

    add_card(slide7, 0.8, 4.3, 5.7, 2.5, "Entitas Lamaran & ATS (Applications)", [
        "Application: UUID (PK), job_id (FK), applicant_id (FK), resume, cover_note, status (Terkirim, Review, Interview, Diterima, Ditolak).",
        "Constraint: Unique Together (job_id, applicant_id) -> Mencegah duplikasi.",
        "ApplicationStatusHistory: application_id (FK), old_status, new_status, notes, changed_by (FK), interview_schedule."
    ], "APLIKASI & ATS", RGBColor(14, 165, 233))

    add_card(slide7, 6.8, 4.3, 5.7, 2.5, "Entitas Komunikasi (Chat)", [
        "Conversation: UUID (PK), applicant_id (FK), company_id (FK), job_id (FK, opsional), updated_at.",
        "ChatMessage: UUID (PK), conversation_id (FK), sender_id (FK), message (text), is_read, created_at.",
        "Relasi: 1 Percakapan dapat menautkan 1 pelamar dengan 1 perusahaan dan konteks lowongan terkait."
    ], "CHAT LANGSUNG", RGBColor(225, 29, 72))

    # ==========================================
    # SLIDE 8: ALUR PENGGUNA (USER JOURNEY & FLOW)
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "7. Alur Kerja Pengguna (User Flow & Recruitment Lifecycle)")

    flow_w = 5.7
    add_card(slide8, 0.8, 1.6, flow_w, 5.2, "Alur Pelamar (Siswa SMK / Pencari Kerja)", [
        "1. Registrasi & Pilih Tipe: Siswa PKL, Mahasiswa, atau Job Seeker.",
        "2. Lengkapi Profil: Masukkan asal sekolah, jurusan SMK, unggah CV & Surat Pengantar PKL.",
        "3. Live Filter Lowongan: Filter berdasarkan kategori PKL SMK, durasi & jurusan yang cocok.",
        "4. Tanya HRD (Opsional): Klik '💬 Chat HRD' untuk menanyakan kuota penempatan.",
        "5. Lamar 1-Klik: Buka modal dialog lamaran, pilih berkas tersimpan, kirim via HTMX.",
        "6. Pantau Lamaran Saya: Cek riwayat status seleksi dan tiket jadwal wawancara terbit."
    ], "FLOW PELAMAR", ACCENT_TEAL)

    add_card(slide8, 6.8, 1.6, flow_w, 5.2, "Alur Perekrut (HRD Perusahaan Mitra)", [
        "1. Registrasi Akun Perusahaan: Verifikasi nama PT, website, dan kebijakan penerimaan PKL.",
        "2. Personalisasi Profil: Unggah logo, banner gedung kantor, dan foto galeri lab/praktik.",
        "3. Posting Lowongan: Tentukan tipe (PKL/Magang/Kerja), jurusan SMK target, dan uang saku.",
        "4. Kelola ATS (Kanban): Review berkas lamaran masuk dan resume PDF kandidat.",
        "5. Update Status Instan: Geser status pelamar ke 'Wawancara' dan masukkan jadwal pertemuan.",
        "6. Chat Langsung: Komunikasi realtime dengan pelamar melalui kotak pesan masuk."
    ], "FLOW PEREKRUT", PRIMARY)

    # ==========================================
    # SLIDE 9: DESIGN SYSTEM GOOGLE MATERIAL DESIGN 3
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "8. Desain Sistem & User Experience (Google Material 3)")

    m_w = 3.65
    add_card(slide9, 0.8, 1.6, m_w, 5.2, "Dynamic Color & Elevation", [
        "Primary Container: Indigo (#4338CA) untuk aksi dominan & navbar.",
        "Tertiary / Accent: Teal (#0D9488) khusus branding penempatan PKL SMK.",
        "Surface Containers: Level 0 s/d Level 4 untuk hierarki visual modul & kartu lowongan.",
        "State Layer: Efek hover dan ripple interaktif khas Google M3."
    ], "SISTEM WARNA M3", PRIMARY)

    add_card(slide9, 4.8, 1.6, m_w, 5.2, "Komponen Standar M3", [
        "Top App Bar: Navigasi responsif dengan profil avatar dropdown & drawer mobile.",
        "Filter Chips & Assist Chips: Tombol seleksi cepat tipe kesempatan & jurusan.",
        "Elevated Cards: Kartu lowongan dengan badge reputasi bintang perusahaan.",
        "Snackbars: Toast alert notifikasi sukses/gagal di pojok kanan bawah."
    ], "KOMPONEN UI", ACCENT_CYAN)

    add_card(slide9, 8.8, 1.6, m_w, 5.2, "Aksesibilitas & Responsivitas", [
        "Mobile-First Experience: Dioptimasi sempurna pada smartphone layar 360px.",
        "Typographic Hierarchy: Skala font Roboto/Inter yang mudah dibaca siswa dan HRD.",
        "Micro-Animations: Transisi modal dan accordion halus dengan Alpine.js.",
        "Zero Layout Shift (CLS): Rendering instan tanpa pergeseran halaman berlebih."
    ], "PENGALAMAN PENGGUNA", ACCENT_TEAL)

    # ==========================================
    # SLIDE 10: ROADMAP & KESIMPULAN (CLOSING)
    # ==========================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    bg10 = slide10.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg10.fill.solid()
    bg10.fill.fore_color.rgb = BG_DARK
    bg10.line.fill.background()

    # Title on Slide 10
    tbox_end = slide10.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(1.0))
    p_e = tbox_end.text_frame.paragraphs[0]
    p_e.text = "9. Rencana Pengembangan (Roadmap) & Penutup"
    p_e.font.size = Pt(24)
    p_e.font.bold = True
    p_e.font.color.rgb = TEXT_WHITE

    # Roadmap Cards (Dark styled)
    def add_dark_card(x, y, w, h, title, items, phase_text, border_col):
        c = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        c.fill.solid()
        c.fill.fore_color.rgb = RGBColor(30, 41, 59)
        c.line.color.rgb = border_col
        c.line.width = Pt(1.5)

        # phase badge
        bdg = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.3), Inches(y + 0.25), Inches(1.8), Inches(0.32))
        bdg.fill.solid()
        bdg.fill.fore_color.rgb = border_col
        bdg.line.fill.background()
        p_b = bdg.text_frame.paragraphs[0]
        p_b.text = phase_text
        p_b.font.size = Pt(9)
        p_b.font.bold = True
        p_b.font.color.rgb = TEXT_WHITE
        p_b.alignment = PP_ALIGN.CENTER

        tb = slide10.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.7), Inches(w - 0.5), Inches(h - 0.8))
        tf = tb.text_frame
        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE

        for itm in items:
            p = tf.add_paragraph()
            p.text = f"• {itm}"
            p.font.size = Pt(11)
            p.font.color.rgb = RGBColor(203, 213, 225)
            p.space_after = Pt(4)

    r_w = 3.65
    add_dark_card(0.8, 2.0, r_w, 4.5, "Fase 1: Rilis MVP Plus (Selesai)", [
        "Pondasi DATH Stack lengkap.",
        "3 Pilar Lowongan (PKL, Magang, Kerja).",
        "Sistem ATS & One-Click Apply.",
        "Fitur Chat In-App Interaktif.",
        "Database Seeding & Rating Perusahaan."
    ], "TAHAP 1 (SELESAI)", ACCENT_TEAL)

    add_dark_card(4.8, 2.0, r_w, 4.5, "Fase 2: Integrasi & Otomasi", [
        "Integrasi Notifikasi WhatsApp Gateway untuk jadwal wawancara.",
        "Tanda Tangan Digital Nota Dinas PKL dari pihak sekolah/SMK.",
        "Sistem Penilaian Akhir Magang & e-Sertifikat otomatis.",
        "Dashboard Evaluasi Guru Pembimbing PKL."
    ], "TAHAP 2 (RENCANA)", ACCENT_CYAN)

    add_dark_card(8.8, 2.0, r_w, 4.5, "Fase 3: Smart AI Matching", [
        "AI Semantic Matching: Mencocokkan CV/Kompetensi siswa dengan lowongan.",
        "Auto-Generator Rekomendasi Jurusan SMK.",
        "Pusat Analitik Serapan Lulusan Vokasi ke Dunia Usaha & Industri (DUDI)."
    ], "TAHAP 3 (SKALABILITAS)", PRIMARY_LIGHT)

    output_path = os.path.join(os.path.dirname(__file__), "Rekrut_Mudah_Presentasi.pptx")
    prs.save(output_path)
    print(f"Presentation successfully created at: {output_path}")

if __name__ == "__main__":
    create_presentation()
